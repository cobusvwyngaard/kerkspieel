#!/usr/bin/env python3
"""
Aggregate the ABR ministers registers into publishable counts.

The raw ABR files list every minister by name, and the 2015-2019 and 2023
files carry South African ID numbers as well. None of that may leave this
script: it reads the registers, derives an age, and writes only counts.

Ages come from two places. Most years carry an OUDERDOM column, but 2020,
2021 and 2022 have it blank for every row, as they do the ID column. The
2022 detail extract (ministers_detail.xlsx) gives a birth year per ABR
number, and ABR numbers are stable between years, so joining on it
recovers those three years at the same ~80% coverage the other years
reach on their own column.

Age is published only for a ring, a synod or the whole church -- never per
congregation, where a single minister's age band would identify them.
Category counts are published per congregation, since how many ministers
a congregation has is not personal information. The output therefore goes
straight to web/public/data, unlike the survey pipeline's intermediate.

Usage:  python3 scripts/build_ministers.py
"""

import argparse
import collections
import csv
import json
import pathlib
import re
import sys

import openpyxl

# Ministers younger than this have not finished the training; older than
# this is a data error (the later files contain ages in the 900s).
AGE_MIN, AGE_MAX = 20, 105
AGE_BANDS = [("<30", 0, 29), ("30-39", 30, 39), ("40-49", 40, 49),
             ("50-59", 50, 59), ("60-64", 60, 64), ("65+", 65, 999)]
# The register uses a "ZZ-" prefix for a minister attached to no
# congregation, and mirrors the credential code into the suffix
# ("ZZ-C02" for an emeritus, "ZZ-B03" for a proponent in no post). They
# are placeholders, not congregations that failed to match.
NO_CONGREGATION_PREFIX = "ZZ-"
UNKNOWN_CATEGORY = "ZZZ"
NATIONAL = "__all__"

# The register groups credentials by the letter of the code. The letter is
# what most questions are actually about; the numbered codes are detail.
CATEGORY_GROUPS = {
    "A": "Predikante en ander volle bevoegdheid",
    "B": "Proponente",
    "C": "Emeriti",
    "D": "Leraars en diensleraars",
    "E": "Gelegitimeer, nie in gemeentebediening",
    "F": "Oorlede",
    "V": "VGK predikante",
}


def text(value):
    return str(value).strip() if value is not None else ""


def load_rows(path, sheet=None, header_row=1):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    rows = [r for r in ws.iter_rows(min_row=header_row + 1, values_only=True)
            if r and any(v not in (None, "") for v in r)]
    wb.close()
    return rows


def band_of(age):
    for key, low, high in AGE_BANDS:
        if low <= age <= high:
            return key
    return None


def read_birth_years(path):
    """ABR number -> birth year, from the 2022 detail extract.

    Only the birth year is taken. The extract also holds names, dates of
    birth and gender; gender is read alongside, nothing else.
    """
    years, gender = {}, {}
    for r in load_rows(path, "20220628_103056"):
        key = text(r[0])
        if not key:
            continue
        try:
            years[key] = int(text(r[15]))
        except ValueError:
            pass
        g = text(r[6]).lower()
        if g.startswith("man"):
            gender[key] = "M"
        elif g.startswith("vrou"):
            gender[key] = "F"
    return years, gender


def read_register(data_dir):
    """Congregation code -> (ring, synod), from the survey register."""
    wb = openpyxl.load_workbook(data_dir / "KS_Ring_Lookup.xlsx",
                                read_only=True, data_only=True)
    ring_rows = [r for r in wb["Sheet1"].iter_rows(min_row=2, values_only=True) if r[1]]
    wb.close()
    wb = openpyxl.load_workbook(data_dir / "sinode_lookup.xlsx",
                                read_only=True, data_only=True)
    synods = {r[0]: r[1] for r in wb["Sheet1"].iter_rows(values_only=True) if r[0]}
    wb.close()

    register = {}
    for ring, code, name in ((r[0], text(r[1]), r[2]) for r in ring_rows):
        prefix = code.split("-")[0]
        if prefix not in synods:
            continue
        ring = text(ring)
        register.setdefault(code, {
            "name": text(name),
            "ring": "" if ring.startswith("#") else ring,
            "synod": synods[prefix],
        })

    # The ABR ring lookup covers congregations the survey register leaves
    # with no ring, so fill from it where it can.
    abr_path = data_dir / "ring_lookup_abr.xlsx"
    if abr_path.exists():
        for r in load_rows(abr_path, "ring lookup"):
            code, ring = text(r[0]), text(r[2])
            if code in register and not register[code]["ring"] and ring:
                register[code]["ring"] = ring
    return register


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--data-dir", default="data/raw")
    ap.add_argument("--out-dir", default="web/public/data")
    args = ap.parse_args()
    data_dir = pathlib.Path(args.data_dir)
    out_dir = pathlib.Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    abr_dir = data_dir / "abr"
    if not abr_dir.exists():
        sys.exit(f"missing {abr_dir}; put the ABR year files there")

    birth_years, gender_of = read_birth_years(data_dir / "ministers_detail.xlsx")
    register = read_register(data_dir)
    print(f"birth years known for {len(birth_years)} ABR numbers")
    print(f"congregation register: {len(register)} congregations")

    labels = {UNKNOWN_CATEGORY: "Onbekend"}
    per_scope = collections.defaultdict(lambda: collections.defaultdict(
        lambda: {"total": 0, "withAge": 0, "ageSum": 0,
                 "categories": collections.Counter(),
                 "groups": collections.Counter(),
                 "bands": collections.Counter(),
                 "gender": collections.Counter()}))
    per_congregation = collections.defaultdict(lambda: collections.defaultdict(
        lambda: {"total": 0, "categories": collections.Counter()}))
    quality = []

    for path in sorted(abr_dir.glob("*.xlsx")):
        year = int(path.stem[:4])
        rows = load_rows(path)
        seen, dropped_age, unknown_cong, dup = set(), 0, 0, 0
        from_column, from_birth = 0, 0

        for r in rows:
            abr = text(r[0])
            # A few later files repeat a minister; count each once.
            if abr and abr in seen:
                dup += 1
                continue
            if abr:
                seen.add(abr)

            code = text(r[3]).upper()
            if not re.fullmatch(r"[A-Z]\d{2}", code):
                # A few files leak a header row; anything else with no usable
                # credential code is a real minister whose category is
                # unrecorded, and is counted as unknown rather than dropped.
                if code in ("DS_STATUS__KODE", "VBO_KODE"):
                    continue
                code = UNKNOWN_CATEGORY
            if text(r[4]):
                labels.setdefault(code, text(r[4]))

            age = None
            raw = text(r[6])
            try:
                value = int(float(raw))
                if AGE_MIN <= value <= AGE_MAX:
                    age = value
                    from_column += 1
                elif raw:
                    dropped_age += 1
            except ValueError:
                pass
            if age is None and abr in birth_years:
                value = year - birth_years[abr]
                if AGE_MIN <= value <= AGE_MAX:
                    age = value
                    from_birth += 1

            congregation = text(r[7])
            attached = bool(congregation) and not congregation.startswith(
                NO_CONGREGATION_PREFIX)
            entry = register.get(congregation) if attached else None
            if attached and not entry:
                # Historic or other-church codes (WT-, VGK...) that the survey
                # register does not carry. They still count nationally.
                unknown_cong += 1

            scopes = [NATIONAL]
            if entry:
                scopes.append(f"s:{entry['synod']}")
                if entry["ring"]:
                    scopes.append(f"r:{entry['synod']}|{entry['ring']}")

            for key in scopes:
                bucket = per_scope[key][year]
                bucket["total"] += 1
                bucket["categories"][code] += 1
                bucket["groups"][code[0]] += 1
                if age is not None:
                    bucket["withAge"] += 1
                    bucket["ageSum"] += age
                    bucket["bands"][band_of(age)] += 1
                if abr in gender_of:
                    bucket["gender"][gender_of[abr]] += 1

            if entry:
                cong = per_congregation[congregation][year]
                cong["total"] += 1
                cong["categories"][code] += 1

        withage = per_scope[NATIONAL][year]["withAge"]
        total = per_scope[NATIONAL][year]["total"]
        quality.append({"year": year, "rows": len(rows), "counted": total,
                        "duplicates": dup, "withAge": withage,
                        "agePct": round(100 * withage / total, 1) if total else 0,
                        "ageOutOfRange": dropped_age,
                        "congregationNotInRegister": unknown_cong,
                        # Where the ages came from. 2020-2022 are wholly
                        # derived, and derived ages only cover ministers who
                        # also appear in the 2022 detail extract, which is a
                        # slightly older group -- so the mean jumps at the
                        # boundary for reasons that are not real ageing.
                        "ageFromColumn": from_column,
                        "ageFromBirthYear": from_birth,
                        "ageDerived": from_column == 0 and from_birth > 0})

    years = sorted({q["year"] for q in quality})
    categories = [{"code": c, "label": labels.get(c, ""), "group": c[0]}
                  for c in sorted(labels)]

    def pack(bucket):
        out = {"total": bucket["total"], "withAge": bucket["withAge"],
               "categories": dict(bucket["categories"]),
               "groups": dict(bucket["groups"])}
        if bucket["withAge"]:
            out["meanAge"] = round(bucket["ageSum"] / bucket["withAge"], 1)
            out["bands"] = [bucket["bands"].get(k, 0) for k, _, _ in AGE_BANDS]
        if bucket["gender"]:
            out["gender"] = dict(bucket["gender"])
        return out

    payload = {
        "years": years,
        "categories": categories,
        "categoryGroups": CATEGORY_GROUPS,
        "ageBands": [k for k, _, _ in AGE_BANDS],
        "ageRange": [AGE_MIN, AGE_MAX],
        # Age is deliberately absent from byCongregation: with one or two
        # ministers, an age band names a person.
        "byScope": {scope: {str(y): pack(b) for y, b in sorted(years_.items())}
                    for scope, years_ in per_scope.items()},
        "byCongregation": {code: {str(y): {"total": b["total"],
                                           "categories": dict(b["categories"])}
                                  for y, b in sorted(years_.items())}
                           for code, years_ in per_congregation.items()},
        # Name, ring and synod for the congregations that appear above, so a
        # report can label and filter them without loading another file.
        "congregations": {code: {"name": register[code]["name"],
                                 "ring": register[code]["ring"],
                                 "synod": register[code]["synod"]}
                          for code in per_congregation},
        "quality": quality,
    }

    path = out_dir / "ministers.json"
    path.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")),
                    encoding="utf-8")

    print()
    print(f"{'year':6}{'rows':>7}{'counted':>9}{'dupes':>7}{'withAge':>9}"
          f"{'age%':>7}{'badAge':>8}{'noCong':>8}")
    for q in quality:
        print(f"{q['year']:<6}{q['rows']:7d}{q['counted']:9d}{q['duplicates']:7d}"
              f"{q['withAge']:9d}{q['agePct']:6.0f}%{q['ageOutOfRange']:8d}"
              f"{q['congregationNotInRegister']:8d}")
    print()
    print(f"scopes: {len(payload['byScope'])}  "
          f"congregations: {len(payload['byCongregation'])}  "
          f"categories: {len(categories)}")
    print(f"wrote {path} ({path.stat().st_size / 1e6:.2f} MB)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
