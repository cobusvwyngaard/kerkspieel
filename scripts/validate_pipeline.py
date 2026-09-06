#!/usr/bin/env python3
"""
Validation harness for the Kerkspieel dataset.

Reproduces figures from the published Power BI "Kerkspieel Ringsverslag"
(WK / George) straight from the raw workbooks, to prove that the
lookup-driven decode path is correct before any dashboard is built.

It also reports the structural defects found in the raw files.

Usage:  python3 scripts/validate_pipeline.py [--data-dir data/raw]
"""

import argparse
import collections
import pathlib
import re
import sys

import openpyxl

SHEETS = {"2018": "KS2018", "2022": "Sheet1"}


def load(path, sheet):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    header = [h.strip() if isinstance(h, str) else h for h in rows[0]]
    data = [r for r in rows[1:] if any(v not in (None, "") for v in r)]
    return header, data


def index(header):
    """Column name -> position. Later duplicates are kept separately."""
    idx = {}
    dupes = collections.defaultdict(list)
    for i, h in enumerate(header):
        if not h:
            continue
        dupes[h].append(i)
        idx.setdefault(h, i)
    return idx, {h: pos for h, pos in dupes.items() if len(pos) > 1}


def read_lookup(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    qs = []
    for r in wb["Sheet1"].iter_rows(min_row=2, values_only=True):
        if not any(v not in (None, "") for v in r):
            continue
        qs.append({"v2018": r[1], "v2022": r[3], "label": r[4], "type": (r[5] or "").lower()})
    opts = collections.defaultdict(dict)
    for r in wb["Sheet2"].iter_rows(min_row=2, values_only=True):
        if not any(v not in (None, "") for v in r):
            continue
        opts[str(r[0]).lower()][r[1]] = {"label": r[2], "value": r[3]}
    wb.close()
    return qs, dict(opts)


def distribution(data, idx, col, options, key_pred):
    """Percentage distribution of a categorical column after option recoding."""
    pos = idx[col]
    kpos = idx["V03"]
    recode = {k: v["value"] for k, v in options.items()}
    vals = [
        recode[r[pos]]
        for r in data
        if key_pred(r[kpos]) and r[pos] in recode
    ]
    total = len(vals)
    counts = collections.Counter(vals)
    return {k: 100 * counts[k] / total for k in sorted(counts)}, total


# Figures transcribed from the published WK / George Ringsverslag PDF.
EXPECTED = [
    ("Huidige rigting", "V160", "V179", "4opsies1",
     {1: 12.11, 2: 49.47, 3: 32.11, 4: 6.32},
     {1: 11.9, 2: 51.4, 3: 31.9, 4: 4.9}),
    ("Finansiele posisie", "V111", "V117", "4opsies2",
     {1: 17.32, 2: 46.93, 3: 29.61, 4: 6.15},
     {1: 18.6, 2: 46.3, 3: 25.5, 4: 9.6}),
    ("Missionale gemeente", "V191", "V225", "lickert5",
     {1: 3.16, 2: 10.00, 3: 12.11, 4: 28.95, 5: 45.79},
     {1: 2.7, 2: 11.4, 3: 10.3, 4: 41.1, 5: 34.6}),
    ("Gemeenskapsbetrokkenheid", "V339", "V378", "4opsies3",
     {1: 3.03, 2: 61.82, 3: 29.70, 4: 5.45},
     {1: 5.4, 2: 58.9, 3: 31.4, 4: 4.3}),
]

TOLERANCE = 0.06  # PDF rounds 2022 figures to one decimal


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--data-dir", default="data/raw")
    args = ap.parse_args()
    d = pathlib.Path(args.data_dir)

    files = {
        "2018": d / "Gemdata_2018_finaal.xlsx",
        "2022": d / "Gemdata_2022_finaal.xlsx",
        "lookup": d / "KS_lookup.xlsx",
        "ring": d / "KS_Ring_Lookup.xlsx",
        "sinode": d / "sinode_lookup.xlsx",
    }
    missing = [str(p) for p in files.values() if not p.exists()]
    if missing:
        sys.exit("missing input files:\n  " + "\n  ".join(missing))

    hdr, data, idx, dup = {}, {}, {}, {}
    for year in ("2018", "2022"):
        hdr[year], data[year] = load(files[year], SHEETS[year])
        idx[year], dup[year] = index(hdr[year])

    questions, options = read_lookup(files["lookup"])

    print("=" * 72)
    print("A. Reproduce published WK-synod figures from raw data")
    print("=" * 72)
    wk = lambda k: isinstance(k, str) and k.startswith("WK")
    failures = 0
    for label, c18, c22, otype, exp18, exp22 in EXPECTED:
        for year, col, exp in (("2018", c18, exp18), ("2022", c22, exp22)):
            got, n = distribution(data[year], idx[year], col, options[otype], wk)
            worst = max(abs(got.get(k, 0) - v) for k, v in exp.items())
            ok = worst <= TOLERANCE
            failures += not ok
            shown = " ".join(f"{k}:{got.get(k, 0):.2f}" for k in sorted(exp))
            print(f"  [{'OK ' if ok else 'FAIL'}] {year} {col:6s} n={n:3d} {label:26s} {shown}")

    print()
    print("=" * 72)
    print("B. Structural defects in the raw files")
    print("=" * 72)
    for year in ("2018", "2022"):
        raw = [h for h in load(files[year], SHEETS[year])[0] if h]
        untrimmed = [h for h in
                     openpyxl.load_workbook(files[year], read_only=True, data_only=True)
                     [SHEETS[year]].iter_rows(max_row=1, values_only=True).__next__()
                     if isinstance(h, str) and h != h.strip()]
        print(f"  {year} headers with stray whitespace: {untrimmed or 'none'}")
        print(f"  {year} duplicate header names       : "
              f"{ {h: p for h, p in dup[year].items()} or 'none'}")

    for year, key in (("2018", "v2018"), ("2022", "v2022")):
        refs = {q[key] for q in questions if q[key]}
        absent = sorted(r for r in refs if r not in idx[year])
        print(f"  lookup refs absent from {year} file : {absent or 'none'}")

    print()
    print("=" * 72)
    print("C. Coverage")
    print("=" * 72)
    for year, key in (("2018", "v2018"), ("2022", "v2022")):
        cols = len([h for h in hdr[year] if h])
        mapped = len({q[key] for q in questions if q[key]} & set(idx[year]))
        print(f"  {year}: {mapped:3d} of {cols} columns carry a label  ({100*mapped/cols:.0f}%)")
    print(f"  questions present in both years          : "
          f"{sum(1 for q in questions if q['v2018'] and q['v2022'])}")
    print(f"  questions new in 2022                    : "
          f"{sum(1 for q in questions if not q['v2018'])}")

    print()
    print("=" * 72)
    print("D. Congregations and participation")
    print("=" * 72)
    hr, dr = load(files["ring"], "Sheet1")
    ring = {r[1]: r[0] for r in dr}
    for year in ("2018", "2022"):
        keys = [r[idx[year]["V03"]] for r in data[year]]
        skip = 4 if year == "2018" else 8
        responded = sum(
            1 for r in data[year]
            if sum(1 for v in r[skip:] if v not in (None, "")) > 20
        )
        print(f"  {year}: {len(keys)} congregations, {responded} responded "
              f"({100*responded/len(keys):.0f}%), "
              f"{sum(1 for k in keys if k not in ring)} not in ring lookup")

    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
