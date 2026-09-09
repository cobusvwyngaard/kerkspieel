#!/usr/bin/env python3
"""
Match questions across waves so a variable can be followed through time.

Codes are renumbered every wave, so comparing 2018 with 2022 with 2026
needs an explicit crosswalk. Matching on question text alone gets 67% of
KS_lookup's verified pairs right: it reliably picks the right question but
then picks the wrong row inside it, because the rows of a grid ("Belydende
lidmate / Dooplidmate / Ongedooptes", "Mans / Vrouens") are worded almost
identically.

So this matches in two stages. Questions are matched first, on their stem,
where the wording is distinctive. The rows inside a matched question are
then aligned in order -- both questionnaires list a grid's rows the same
way -- falling back to row-text similarity when the counts differ.

KS_lookup is treated as ground truth: its pairs are asserted first and
also used to score the rest.

Usage:  python3 scripts/build_crosswalk.py
"""

import argparse
import collections
import csv
import difflib
import json
import pathlib
import re
import sys
import unicodedata

import openpyxl

WAVES = ("2018", "2022", "2026")
QUESTION_THRESHOLD = 0.55


def words(text):
    text = unicodedata.normalize("NFKD", str(text or ""))
    text = text.encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9 ]+", " ", text).split()


def similarity(a, b):
    """Blend token overlap with sequence ratio; neither alone is enough."""
    wa, wb = set(words(a)), set(words(b))
    if not wa or not wb:
        return 0.0
    jaccard = len(wa & wb) / len(wa | wb)
    ratio = difflib.SequenceMatcher(None, " ".join(words(a)), " ".join(words(b))).ratio()
    return 0.5 * jaccard + 0.5 * ratio


def load_codebook(path):
    """code -> variable, plus the questions in sheet order."""
    raw = json.loads(path.read_text(encoding="utf-8"))
    by_position = {int(j): v for j, v in raw["resolved"].items()}
    by_code, questions = {}, collections.OrderedDict()
    for j in sorted(by_position):
        v = dict(by_position[j], position=j)
        by_code[v["code"]] = v
        questions.setdefault((v["qnum"], v["stem"]), []).append(v)
    return by_code, questions


def read_truth(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    pairs = [(r[1], r[3], r[4], (r[5] or "").lower())
             for r in wb["Sheet1"].iter_rows(min_row=2, values_only=True)
             if r and (r[1] or r[3])]
    wb.close()
    return pairs


def match_waves(src_questions, dst_questions):
    """Pair codes between two waves, question block first then row.

    Blocks are assigned best-scoring-pair-first across the whole grid
    rather than in source order: taking each source question in turn lets
    an early mediocre match consume the block a later question needed.
    """
    scored = sorted(
        ((similarity(src_key[1], dst_key[1]), src_key, dst_key)
         for src_key in src_questions for dst_key in dst_questions),
        key=lambda t: -t[0])

    matched, used_src, used_dst, pairs = [], set(), set(), {}
    for score, src_key, dst_key in scored:
        if score < QUESTION_THRESHOLD:
            break
        if src_key in used_src or dst_key in used_dst:
            continue
        used_src.add(src_key)
        used_dst.add(dst_key)
        matched.append((score, src_key, dst_key))

    for best_score, src_key, best_key in matched:
        src_rows = src_questions[src_key]
        dst_rows = dst_questions[best_key]

        if len(src_rows) == len(dst_rows):
            # A grid keeps its rows in the same order between waves.
            for s, d in zip(src_rows, dst_rows):
                pairs[s["code"]] = (d["code"], round(best_score, 3), "block+order")
        else:
            taken = set()
            for s in src_rows:
                cands = [d for d in dst_rows if d["code"] not in taken]
                if not cands:
                    continue
                d = max(cands, key=lambda d: similarity(
                    s.get("item") or s["stem"], d.get("item") or d["stem"]))
                taken.add(d["code"])
                pairs[s["code"]] = (d["code"], round(best_score, 3), "block+text")
    return pairs


def with_scales(row, scale):
    """Record how many options each wave offered, and whether that held.

    A question can survive a redesign and still not be comparable: the
    youth-ministry grids keep their wording but drop from four options in
    2022 to three in 2026, so charting the waves together would be wrong.
    """
    sizes = {y: scale(y, row[f"V{y}"], row.get("type", "")) if row[f"V{y}"] else 0
             for y in WAVES}
    present = [n for n in sizes.values() if n]
    row.update({f"options_{y}": sizes[y] or "" for y in WAVES})
    waves_present = sum(1 for y in WAVES if row[f"V{y}"])
    if waves_present < 2:
        row["comparable"] = "single-wave"
    elif not present:
        row["comparable"] = "numeric"
    elif len(set(present)) == 1 and len(present) == waves_present:
        row["comparable"] = "yes"
    else:
        row["comparable"] = "scale-changed"
    return row


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--codebook-dir", default="data/codebook")
    ap.add_argument("--data-dir", default="data/raw")
    ap.add_argument("--out", default="data/lookup/crosswalk.csv")
    args = ap.parse_args()

    cb_dir = pathlib.Path(args.codebook_dir)
    books = {}
    for year in WAVES:
        path = cb_dir / f"codebook_{year}.json"
        if not path.exists():
            sys.exit(f"missing {path}; run scripts/build_codebook.py first")
        books[year] = load_codebook(path)

    options_path = cb_dir / "options.json"
    options = json.loads(options_path.read_text(encoding="utf-8")) \
        if options_path.exists() else {y: {} for y in WAVES}

    # The Likert grids name their columns in a header well above the rows,
    # which the questionnaire parser does not reach; KS_lookup names those
    # scales outright, so fall back to it before calling a question numeric.
    wb = openpyxl.load_workbook(pathlib.Path(args.data_dir) / "KS_lookup.xlsx",
                                read_only=True, data_only=True)
    typed_size = collections.Counter(str(r[0]).lower()
                                     for r in wb["Sheet2"].iter_rows(min_row=2,
                                                                     values_only=True)
                                     if r and r[0])
    wb.close()

    def scale(year, code, kind=""):
        """Number of response options, or 0 where the question is numeric."""
        listed = options.get(year, {}).get(code)
        return len(listed) if listed else typed_size.get(kind.lower(), 0)

    truth = read_truth(pathlib.Path(args.data_dir) / "KS_lookup.xlsx")
    verified = {a: b for a, b, _, _ in truth if a and b}
    # Keyed by wave, never by "whichever code the row has". KS_lookup holds
    # rows that start at 2022 -- the hospitality grid is one -- and a single
    # dict keyed by either code lets a 2022 code stand in for a 2018 one.
    # V241 exists in both waves and names a different question in each, so
    # that collision silently gave seven questions someone else's label.
    labels = {wave: {} for wave in ("2018", "2022")}
    types = {wave: {} for wave in ("2018", "2022")}
    for a, b, lab, kind in truth:
        for wave, code in (("2018", a), ("2022", b)):
            if code:
                labels[wave][code] = lab
                types[wave][code] = kind

    derived = match_waves(books["2018"][1], books["2022"][1])
    agree = sum(1 for a, b in verified.items() if derived.get(a, ("",))[0] == b)
    checked = sum(1 for a in verified if a in derived)
    print(f"2018->2022 against KS_lookup's {len(verified)} verified pairs: "
          f"{agree}/{checked} agree ({100 * agree / max(checked, 1):.0f}%)")

    forward = match_waves(books["2022"][1], books["2026"][1])
    print(f"2022->2026 derived pairs: {len(forward)}")

    rows = []
    for code18, v18 in books["2018"][0].items():
        code22 = verified.get(code18) or derived.get(code18, (None,))[0]
        source = "KS_lookup" if code18 in verified else (
            derived[code18][2] if code18 in derived else "")
        code26 = forward.get(code22, (None,))[0] if code22 else None
        rows.append(with_scales({
            "label": labels["2018"].get(code18)
                     or (v18["stem"] + " " + v18.get("item", "")).strip(),
            "type": types["2018"].get(code18, ""),
            "V2018": code18, "V2022": code22 or "", "V2026": code26 or "",
            "source_2018_2022": source,
            "source_2022_2026": forward.get(code22, ("", "", ""))[2] if code22 else "",
        }, scale))

    # Questions that only appear later still belong in the crosswalk.
    seen22 = {r["V2022"] for r in rows if r["V2022"]}
    for code22, v22 in books["2022"][0].items():
        if code22 in seen22:
            continue
        code26 = forward.get(code22, (None,))[0]
        rows.append(with_scales({
            "label": labels["2022"].get(code22)
                     or (v22["stem"] + " " + v22.get("item", "")).strip(),
            "type": types["2022"].get(code22, ""), "V2018": "", "V2022": code22,
            "V2026": code26 or "", "source_2018_2022": "",
            "source_2022_2026": forward.get(code22, ("", "", ""))[2] if code26 else "",
        }, scale))
    seen26 = {r["V2026"] for r in rows if r["V2026"]}
    for code26, v26 in books["2026"][0].items():
        if code26 in seen26:
            continue
        rows.append(with_scales({"label": v26.get("label") or v26["stem"], "type": "",
                                 "V2018": "", "V2022": "", "V2026": code26,
                                 "source_2018_2022": "", "source_2022_2026": ""}, scale))

    out = pathlib.Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    with out.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=["label", "type", "V2018", "V2022",
                                           "V2026", "options_2018", "options_2022",
                                           "options_2026", "comparable",
                                           "source_2018_2022", "source_2022_2026"])
        w.writeheader()
        w.writerows(rows)

    all_three = sum(1 for r in rows if r["V2018"] and r["V2022"] and r["V2026"])
    verdicts = collections.Counter(r["comparable"] for r in rows)
    print(f"\ncrosswalk rows: {len(rows)}")
    for verdict, n in verdicts.most_common():
        print(f"  {verdict:14s} {n:4d}")
    print(f"  present in all three waves : {all_three}")
    print(f"  2018 and 2022 only         : "
          f"{sum(1 for r in rows if r['V2018'] and r['V2022'] and not r['V2026'])}")
    print(f"  2026 only (new questions)  : "
          f"{sum(1 for r in rows if r['V2026'] and not r['V2022'])}")
    print(f"wrote {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
