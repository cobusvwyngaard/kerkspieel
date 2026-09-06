#!/usr/bin/env python3
"""
Assemble the browser-facing dataset from the raw workbooks.

Everything the dashboard needs is precomputed here and shipped as static
JSON: there is no server and no database, because all three waves together
compress to a fraction of a megabyte.

Output, into web/public/data:
  congregations.json  the standing register: code, name, ring, synod, point
  questions.json      the crosswalk, with each wave's option labels
  responses.json      one row per congregation per wave, values keyed by
                      the crosswalk's question id rather than by a code
                      that means something different in every wave

Usage:  python3 scripts/build_dataset.py
"""

import argparse
import collections
import csv
import json
import pathlib
import re
import sys
import unicodedata

import openpyxl

WAVES = ("2018", "2022", "2026")
SHEETS = {"2018": ("Gemdata_2018_finaal.xlsx", "KS2018", 1),
          "2022": ("Gemdata_2022_finaal.xlsx", "Sheet1", 1),
          "2026": ("CS_2026_NGK_14072026.xlsx",
                   "Sheet 1 - Congregational_survey", 2)}
# Below this many answered columns a row is a non-response, not a sparse one.
ANSWERED_MIN = 20
NOISE = re.compile(r"^(ng\s*kerk|ng\s*gemeente|ngk|n\s*g\s*kerk|n\s*g\s*gemeente|"
                   r"nederduitse\s+gereformeerde\s+kerk|gemeente|ng)\b[\s.\-]*", re.I)


def load_sheet(path, sheet, header_row):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(wb[sheet].iter_rows(values_only=True))
    wb.close()
    return [r for r in rows[header_row:] if any(v not in (None, "") for v in r)]


def normalise_name(value):
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = text.encode("ascii", "ignore").decode().lower().replace("\xa0", " ").strip()
    for _ in range(3):
        text = NOISE.sub("", text).strip()
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def read_csv(path):
    lines = [ln for ln in path.read_text(encoding="utf-8").splitlines()
             if not ln.lstrip().startswith("#")]
    return list(csv.DictReader(lines))


def build_register(data_dir):
    """The standing congregation register, with coordinates where known."""
    wb = openpyxl.load_workbook(data_dir / "KS_Ring_Lookup.xlsx",
                                read_only=True, data_only=True)
    ring_rows = [r for r in wb["Sheet1"].iter_rows(min_row=2, values_only=True) if r[1]]
    wb.close()
    wb = openpyxl.load_workbook(data_dir / "sinode_lookup.xlsx",
                                read_only=True, data_only=True)
    synods = {r[0]: r[1] for r in wb["Sheet1"].iter_rows(values_only=True) if r[0]}
    wb.close()

    points = {}
    for name, status in (("Gemeentes1.xlsx", "active"),
                         ("Ontbinde_gemeentes.xlsx", "dissolved")):
        wb = openpyxl.load_workbook(data_dir / name, read_only=True, data_only=True)
        for r in wb["sql_statement"].iter_rows(min_row=2, values_only=True):
            if r and r[3]:
                points.setdefault(normalise_name(r[3]), (float(r[0]), float(r[1]), status))
        wb.close()

    register = {}
    for ring, code, name in ((r[0], r[1].strip(), r[2]) for r in ring_rows):
        point = points.get(normalise_name(name))
        register.setdefault(code, {
            "code": code, "name": name,
            # The ring lookup carries a few spelling variants of the same
            # ring; the synod comes from the code prefix, which is cleaner
            # than the synod column in either data file.
            "ring": (ring or "").strip(),
            "synod": synods.get(code.split("-")[0], code.split("-")[0]),
            "lon": point[0] if point else None,
            "lat": point[1] if point else None,
        })
    return register


def resolved_keys_2026(lookup_dir):
    """row index -> congregation code, for rows kept after de-duplication."""
    path = lookup_dir / "key_2026_resolved.csv"
    if not path.exists():
        sys.exit(f"missing {path}; run scripts/resolve_2026_keys.py first")
    return {int(r["row"]): r["V03"] for r in read_csv(path)
            if r["V03"] and r.get("status", "keep") == "keep"}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--data-dir", default="data/raw")
    ap.add_argument("--lookup-dir", default="data/lookup")
    ap.add_argument("--codebook-dir", default="data/codebook")
    ap.add_argument("--out-dir", default="web/public/data")
    args = ap.parse_args()

    data_dir = pathlib.Path(args.data_dir)
    lookup_dir = pathlib.Path(args.lookup_dir)
    cb_dir = pathlib.Path(args.codebook_dir)
    out_dir = pathlib.Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    register = build_register(data_dir)
    options = json.loads((cb_dir / "options.json").read_text(encoding="utf-8"))
    crosswalk = read_csv(lookup_dir / "crosswalk.csv")

    # Where a question's position on the scale is known, order the options
    # the way the published report does -- worst first -- rather than by the
    # raw code, which several questions record in the opposite direction.
    wb = openpyxl.load_workbook(data_dir / "KS_lookup.xlsx", read_only=True, data_only=True)
    display_order = collections.defaultdict(dict)
    typed_options = collections.defaultdict(list)
    for r in wb["Sheet2"].iter_rows(min_row=2, values_only=True):
        if r and r[0]:
            kind = str(r[0]).lower()
            display_order[kind][r[1]] = r[3]
            typed_options[kind].append({"value": r[1], "label": r[2]})
    wb.close()

    questions, code_index = [], {y: {} for y in WAVES}
    for n, row in enumerate(crosswalk):
        qid = f"q{n:03d}"
        opts = {}
        for year in WAVES:
            code = row[f"V{year}"]
            if code:
                code_index[year][code] = qid
                listed = options.get(year, {}).get(code)
                # A scale read off trailing digits alone has values but no
                # wording; prefer KS_lookup's names in that case.
                if listed and any(lab for _, lab in listed):
                    opts[year] = [{"value": v, "label": lab} for v, lab in listed]
                elif row["type"].lower() in typed_options:
                    # The Likert grids state their columns in a header several
                    # lines above the rows, which the questionnaire parser does
                    # not reach; KS_lookup names those scales outright.
                    opts[year] = [dict(o) for o in typed_options[row["type"].lower()]]
        questions.append({
            "id": qid, "label": row["label"], "type": row["type"],
            "codes": {y: row[f"V{y}"] for y in WAVES if row[f"V{y}"]},
            "options": opts,
            "order": display_order.get(row["type"].lower(), {}),
            "comparable": row["comparable"],
        })

    keys_2026 = resolved_keys_2026(lookup_dir)
    responses = []
    for year in WAVES:
        filename, sheet, header_row = SHEETS[year]
        rows = load_sheet(data_dir / filename, sheet, header_row)
        book = json.loads((cb_dir / f"codebook_{year}.json").read_text(encoding="utf-8"))
        columns = [(int(j), v["code"]) for j, v in book["resolved"].items()]
        key_col = None if year == "2026" else next(
            int(j) for j, v in book["resolved"].items() if v["code"] == "V03")

        for i, row in enumerate(rows):
            code = keys_2026.get(i) if year == "2026" else row[key_col]
            if not code or code not in register:
                continue
            values, answered = {}, 0
            for j, var in columns:
                value = row[j] if j < len(row) else None
                if value in (None, ""):
                    continue
                answered += 1
                qid = code_index[year].get(var)
                if qid and isinstance(value, (int, float)):
                    values[qid] = value
            if answered >= ANSWERED_MIN:
                responses.append({"wave": year, "code": code, "values": values})

    counts = collections.Counter(r["wave"] for r in responses)
    for name, payload in (("congregations", sorted(register.values(), key=lambda c: c["code"])),
                          ("questions", questions),
                          ("responses", responses)):
        path = out_dir / f"{name}.json"
        path.write_text(json.dumps(payload, ensure_ascii=False,
                                   separators=(",", ":")), encoding="utf-8")
        print(f"  {path}  {path.stat().st_size / 1e6:.2f} MB")

    placed = sum(1 for c in register.values() if c["lat"] is not None)
    print(f"\ncongregations: {len(register)} ({placed} with coordinates)")
    print(f"questions: {len(questions)}")
    print("responses: " + ", ".join(f"{y} {counts[y]}" for y in WAVES))
    return 0


if __name__ == "__main__":
    sys.exit(main())
